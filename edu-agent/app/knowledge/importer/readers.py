"""
多格式文件读取器 - 统一接口支持多种文件格式。

支持格式：
- .md / .txt：原生 Python 读取
- .pdf：pdfplumber 解析
- .docx：python-docx 解析

依赖安装：
    pip install pdfplumber python-docx
"""
import re
from pathlib import Path

from app.common.logging import logger


def read_file(file_path: str) -> str:
    """
    根据文件扩展名自动选择读取方式，返回纯文本。

    Args:
        file_path: 文件路径

    Returns:
        文件的纯文本内容
    """
    path = Path(file_path)
    ext = path.suffix.lower()

    readers = {
        ".md": _read_text,
        ".txt": _read_text,
        ".pdf": _read_pdf,
        ".docx": _read_docx,
    }

    reader = readers.get(ext)
    if reader is None:
        raise ValueError(f"不支持的文件格式: {ext}，支持: {list(readers.keys())}")

    logger.info(f"读取文件: {path.name}（{ext}）")
    return reader(path)


def _read_text(path: Path) -> str:
    """读取纯文本文件（.md / .txt）。"""
    return path.read_text(encoding="utf-8")


def _read_pdf(path: Path) -> str:
    """读取 PDF 文件。"""
    import pdfplumber

    text_parts = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages):
            page_text = page.extract_text()
            if page_text:
                text_parts.append(f"--- 第 {i + 1} 页 ---\n{page_text}")

    if not text_parts:
        logger.warning(f"PDF 文件无文本内容: {path.name}")
        return ""

    return "\n\n".join(text_parts)


def _read_docx(path: Path) -> str:
    """读取 DOCX 文件。"""
    from docx import Document

    doc = Document(path)
    text_parts = []

    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells]
            if any(row_cells):
                text_parts.append(" | ".join(row_cells))

    return "\n".join(text_parts)


def detect_file_type(file_path: str, content: str = "") -> str:
    """
    检测文件类型，返回 'course_intro' / 'questions' / 'generic'。

    两级检测策略（防止文件名关键字误判）：
    1. 文件名提示（弱信号）：含"课程"/"题目"等关键字，标记为候选
    2. 内容结构校验（强信号）：检查是否匹配预定义格式特征

    只有文件名提示 + 内容校验都通过，才走专用解析路径。
    否则一律走 generic 通用路径。

    Args:
        file_path: 文件路径
        content: 文件内容（用于结构校验）

    Returns:
        文件类型字符串
    """
    path = Path(file_path)
    name = path.stem

    has_course_keyword = "课程" in name
    has_question_keyword = "题目" in name or "题库" in name

    if not has_course_keyword and not has_question_keyword:
        return "generic"

    # 文件名有关键字 → 用内容结构校验
    if content:
        if has_course_keyword and _looks_like_course_intro(content):
            logger.info(f"文件 {path.name} 匹配课程介绍结构")
            return "course_intro"

        if has_question_keyword and _looks_like_questions(content):
            logger.info(f"文件 {path.name} 匹配题目资料结构")
            return "questions"

    # 内容校验未通过 → 当通用文档处理
    logger.info(f"文件 {path.name} 关键字匹配但结构不通过，按通用文档处理")
    return "generic"


def _looks_like_course_intro(content: str) -> bool:
    """
    检查内容是否匹配课程介绍的结构特征。

    课程介绍特征：
    - 有 ## 标题行
    - 包含"系列编码"、"课程分类"等结构化字段
    - 有"编码: xxx, 课时:"格式的模块列表
    """
    has_series_code = bool(re.search(r'系列编码.*?:\s*\S+', content))
    has_category = bool(re.search(r'课程分类.*?:\s*', content))
    has_module_coding = bool(re.search(r'编码:\s*\S+?,\s*课时:', content))
    has_hash_heading = bool(re.search(r'^##\s+', content, re.MULTILINE))

    match_count = sum([has_series_code, has_category, has_module_coding, has_hash_heading])
    return match_count >= 2


def _looks_like_questions(content: str) -> bool:
    """
    检查内容是否匹配题目资料的结构特征。

    题目资料特征：
    - 有 ### 标题行（每题一个）
    - 包含"题型"、"题干"、"选项"、"答案"等结构化字段
    - 有 A/B/C/D 选项格式
    """
    has_question_type = bool(re.search(r'题型.*?:\s*', content))
    has_stem = bool(re.search(r'题干.*?:\s*', content))
    has_options = bool(re.search(r'选项.*?:\s*\n', content))
    has_answer = bool(re.search(r'答案.*?:\s*[A-D]$', content, re.MULTILINE))
    has_hash3 = bool(re.search(r'^###\s+', content, re.MULTILINE))

    match_count = sum([has_question_type, has_stem, has_options, has_answer, has_hash3])
    return match_count >= 2
